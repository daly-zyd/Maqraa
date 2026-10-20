from datetime import datetime, time

import openpyxl
import pytest

from apps.people.models import Person, PersonAlias
from apps.scheduling import services
from apps.scheduling.models import WeeklyTemplateSlot
from integrations import planning_xlsx

DAYS = ["الجمعة", "الخميس", "الإربعاء", "الثلاثاء", "الإثنين", "الأحد", "السبت"]  # Hafs sheet order


def build_sheet(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws["C1"] = "مقرأة حفص رجال بالقرآن نحيا"
    ws["S1"] = datetime(2026, 10, 8)
    ws.append(["تونس", "مكة", None, *[x for d in DAYS for x in (d, None)]])
    # header row is row 2 -> day columns D, F, H, ...
    def row(tunis, mecca, label, cells):
        ws.append([tunis, mecca, label, *[cells.get(d) for d in DAYS for _ in (0, 1)]])
    def block(tunis, mecca, sheikhs, sups=None):
        r = [None] * 14
        for d, name in sheikhs.items():
            r[DAYS.index(d) * 2] = name
        ws.append([tunis, mecca, "المصحح", *r])
        r2 = [None] * 14
        for d, name in (sups or {}).items():
            r2[DAYS.index(d) * 2] = name
        ws.append([None, None, "المشرف", *r2])
    block(time(9, 0), time(11, 0), {"الخميس": "ش/محمد سويد", "الأحد": "ش/محمد سويد", "الجمعة": "ش/محمد سويد"},
          {"الخميس": "حاتم بن عياد "})
    block(time(18, 0), time(19, 0), {"الخميس": "ش/ فؤاد حسان"})            # inconsistent row -> 17:00
    block(time(18, 0), time(20, 0), {"الأربعاء".replace("الأ", "الإ"): None, "الإربعاء": "ش/ يوسف المديحي"})
    block(time(5, 0), time(7, 0), {"الثلاثاء": "ش/أحمد بن عياد (05:30)"})
    block(time(19, 0), time(21, 0), {"الخميس": "ش/ يوسف مديحي", "الأحد": "ش/ يوسف مديحي"})
    wb.save(path)


@pytest.fixture
def sheet(tmp_path):
    p = tmp_path / "p.xlsx"
    build_sheet(p)
    return p


def test_parser_reads_days_times_overrides_and_warnings(sheet):
    parsed = planning_xlsx.parse_workbook(sheet)
    assert parsed.maqraa_slug_hint == "hafs-men"
    by_name = {(s.corrector, s.weekday): s for s in parsed.slots}
    thursday = 3
    assert by_name[("ش/محمد سويد", thursday)].start == time(9, 0)
    fouad = by_name[("ش/ فؤاد حسان", thursday)]
    assert fouad.start == time(17, 0)  # Mecca 19:00 - 2h; the Tunis cell said 18:00
    override = by_name[("ش/أحمد بن عياد (05:30)", 1)]
    assert (override.start, override.end) == (time(5, 30), time(6, 30))
    assert by_name[("ش/محمد سويد", thursday)].supervisor == "حاتم بن عياد"
    assert ("ش/محمد سويد", 4) in by_name  # Friday (الجمعة, taa-marbuta) must not be dropped
    assert {s.weekday for s in parsed.slots} >= {1, 3, 4, 6}
    assert any("inconsistent" in w for w in parsed.warnings) and any("override" in w for w in parsed.warnings)


def test_import_creates_people_aliases_roles_and_slots(sheet, make_user, maq):
    boss = make_user("boss", super_admin=True)
    parsed = planning_xlsx.parse_workbook(sheet)
    report = services.import_planning(boss, maq["hafs-men"].pk, parsed)
    assert report["created"] == len(parsed.slots) and not report["skipped"]
    names = set(Person.objects.system().values_list("full_name", flat=True))
    assert "يوسف المديحي" in names and not any(n.startswith("ش/") for n in names)
    # "يوسف المديحي" and "يوسف مديحي" are ONE person
    assert Person.objects.system().filter(name_key="يوسف مديحي").count() == 1
    assert PersonAlias.objects.filter(person__name_key="يوسف مديحي").count() == 1
    assert WeeklyTemplateSlot.objects.system().filter(maqraa=maq["hafs-men"], active=True).count() == len(parsed.slots)


def test_import_is_replace_by_default_and_idempotent(sheet, make_user, maq):
    boss = make_user("boss", super_admin=True)
    parsed = planning_xlsx.parse_workbook(sheet)
    services.import_planning(boss, maq["hafs-men"].pk, parsed)
    services.import_planning(boss, maq["hafs-men"].pk, parsed)
    assert WeeklyTemplateSlot.objects.system().filter(maqraa=maq["hafs-men"], active=True).count() == len(parsed.slots)
    assert Person.objects.system().count() == len({p.pk for p in Person.objects.system().all()})


def test_import_skips_cross_maqraa_conflicts_and_reports_them(sheet, make_user, maq):
    boss = make_user("boss", super_admin=True)
    parsed = planning_xlsx.parse_workbook(sheet)
    services.import_planning(boss, maq["hafs-men"].pk, parsed)
    # the same grid imported into Qaloun men: same people at the same hours -> blocked, not duplicated
    report = services.import_planning(boss, maq["qaloun-men"].pk, parsed)
    assert report["created"] == 0 and len(report["skipped"]) == len(parsed.slots)
    assert {s["reason"] for s in report["skipped"]} == {"person_busy_elsewhere"}


def test_import_requires_admin_in_the_target_maqraa(sheet, make_user, maq):
    from apps.core.exceptions import ScopeNotFound

    admin = make_user("adm", roles={"qaloun-men": ["admin"]})
    with pytest.raises(ScopeNotFound):
        services.import_planning(admin, maq["hafs-men"].pk, planning_xlsx.parse_workbook(sheet))
